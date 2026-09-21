"""Build M2 Session 2 deck: Data Preparation and Feature Engineering.

All numbers come from M2/01-data-prep-and-feature-engineering.ipynb
(executed outputs, see notebook brief). Charts regenerated from those values.
"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx.util import Inches
import deck_style as ds

pal = ds.mpl_theme()


def notes(slide, text):
    """Speaker notes: a plain-text talk track for the teacher."""
    slide.notes_slide.notes_text_frame.text = text


FIGS = "/tmp/deck-workshop/figs2"
FIGS_STORY = "/tmp/deck-workshop/figs-story"
FIGS_T2 = "/tmp/deck-workshop/figs-theory2"
os.makedirs(FIGS, exist_ok=True)
os.makedirs(FIGS_STORY, exist_ok=True)
os.makedirs(FIGS_T2, exist_ok=True)

# ---------------- Figure 1: imputation - fit on train only ----------------
fig, axes = plt.subplots(1, 2, figsize=(7.5, 4))
ax = axes[0]
vals = [137, 40]
ax.bar(["Train\n(712 rows)", "Test\n(179 rows)"], vals,
       color=[pal["blue"], pal["sky"]])
ax.set_title("Missing Age rows")
for i, v in enumerate(vals):
    ax.text(i, v + 4, str(v), ha="center", fontweight="bold", color=pal["ink"])
ax.set_ylim(0, 160)
ax = axes[1]
meds = [28.5, 27.0]
ax.bar(["Train median\n(used for both)", "Test median\n(never used)"], meds,
       color=[pal["blue"], pal["sky"]])
ax.set_title("Median Age learned")
for i, v in enumerate(meds):
    ax.text(i, v + 0.6, f"{v:.1f}", ha="center", fontweight="bold",
            color=pal["ink"])
ax.set_ylim(0, 33)
fig.tight_layout()
fig.savefig(f"{FIGS}/fig_impute.png")
plt.close(fig)

# ---------------- Figure 2: raw scales - Fare dwarfs Age ----------------
fig, ax = plt.subplots(figsize=(7.5, 4))
feats = ["Age", "Fare"]
maxes = [80, 512.3]
labels = ["0 to 80", "0 to 512.3"]
bars = ax.barh(feats, maxes, color=[pal["sky"], pal["blue"]])
for bar, lab in zip(bars, labels):
    ax.text(bar.get_width() + 8, bar.get_y() + bar.get_height() / 2, lab,
            va="center", fontweight="bold", color=pal["ink"])
ax.set_title("Raw value ranges before scaling")
ax.set_xlabel("Value")
ax.set_xlim(0, 610)
fig.tight_layout()
fig.savefig(f"{FIGS}/fig_scales.png")
plt.close(fig)

# ---------------- Figure 3: grouped Title counts (train) ----------------
fig, ax = plt.subplots(figsize=(7.5, 4))
titles = ["Rare", "Master", "Mrs", "Miss", "Mr"]  # bottom-up for barh
counts = [18, 31, 107, 144, 412]
colors = [pal["sky"], pal["blue"], pal["sky"], pal["sky"], pal["sky"]]
bars = ax.barh(titles, counts, color=colors)
for bar, v in zip(bars, counts):
    ax.text(bar.get_width() + 6, bar.get_y() + bar.get_height() / 2, str(v),
            va="center", fontweight="bold", color=pal["ink"])
ax.set_title("Title counts after grouping, training set")
ax.set_xlabel("Passengers")
ax.set_xlim(0, 470)
fig.tight_layout()
fig.savefig(f"{FIGS}/fig_titles.png")
plt.close(fig)

# ---------------- Figure 4: FamilySize distribution (train) ----------------
fig, ax = plt.subplots(figsize=(7.5, 4))
sizes = ["1", "2", "3", "4", "5", "6", "7", "8", "11"]
fam = [434, 129, 77, 22, 14, 16, 11, 4, 5]
colors = [pal["blue"] if s in ("2", "3", "4") else pal["sky"] for s in sizes]
bars = ax.bar(sizes, fam, color=colors)
for bar, v in zip(bars, fam):
    ax.text(bar.get_x() + bar.get_width() / 2, v + 8, str(v), ha="center",
            fontweight="bold", color=pal["ink"], fontsize=11)
ax.set_title("FamilySize distribution, training set (sizes 2-4 highlighted)")
ax.set_xlabel("FamilySize = SibSp + Parch + 1")
ax.set_ylabel("Passengers")
ax.set_ylim(0, 490)
fig.tight_layout()
fig.savefig(f"{FIGS}/fig_family.png")
plt.close(fig)

# -------- Frame C (cross-deck): the FIT | CHOOSE | REPORT verbs bar --------
# One canonical drawing, shared with Decks 3 and 4 via shared_figs.
from matplotlib.patches import Rectangle
from shared_figs import make_frame_c

make_frame_c(f"{FIGS_T2}/frame_c_bar_s2.png",
             tag="the hatched band = the CHOOSE step; we build it in Session 3")

# -------- Theory figure: MCAR / MAR / MNAR three-card diagram --------
CARDS = [
    ("MCAR", "missing completely at random", pal["sky"], pal["navy"], [
        ("MEANING", "Holes land blindly -\nnothing decides where"),
        ("TITANIC STORY (PLAUSIBLE)", "What if a clerk copying the\nmanifest skipped lines at random?"),
        ("IF IGNORED", "Rows lost, but no bias"),
        ("WHAT TO DO", "Drop or impute simply - safe"),
    ]),
    ("MAR", "missing at random", pal["blue"], "white", [
        ("MEANING", "The chance of a hole depends\non OTHER observed columns"),
        ("TITANIC STORY (PLAUSIBLE)", "What if third-class ages were\nrecorded less often than first?"),
        ("IF IGNORED", "Results become biased"),
        ("WHAT TO DO", "Model the fill from the\ncolumns you did observe"),
    ]),
    ("MNAR", "missing not at random", pal["navy"], "white", [
        ("MEANING", "The hole depends on the\nmissing value itself"),
        ("TITANIC STORY (PLAUSIBLE)", "What if an age was kept off the\nrecord because of the age itself?"),
        ("IF IGNORED", "Bias no fill can see"),
        ("WHAT TO DO", "No imputation fully repairs\nit - reason about the world"),
    ]),
]
fig, ax = plt.subplots(figsize=(7.3, 4.5))
for i, (name, sub, head_fc, head_tc, rows) in enumerate(CARDS):
    x = 0.15 + i * 4.0
    ax.add_patch(Rectangle((x, 0.2), 3.7, 8.1, facecolor=pal["panel"],
                           edgecolor=pal["sky"], lw=1.2))
    ax.add_patch(Rectangle((x, 8.3), 3.7, 1.5, facecolor=head_fc))
    ax.text(x + 1.85, 9.25, name, ha="center", va="center", fontsize=13,
            fontweight="bold", color=head_tc)
    ax.text(x + 1.85, 8.62, sub, ha="center", va="center", fontsize=7.2,
            color=head_tc)
    y = 7.75
    for label, text in rows:
        ax.text(x + 0.18, y, label, ha="left", va="top", fontsize=6.4,
                fontweight="bold", color=pal["gray"])
        ax.text(x + 0.18, y - 0.55, text, ha="left", va="top", fontsize=8.2,
                color=pal["ink"])
        y -= 0.62 + 0.62 * (text.count("\n") + 1) + 0.28
ax.set_xlim(0, 12.1)
ax.set_ylim(0, 9.9)
ax.axis("off")
fig.savefig(f"{FIGS_T2}/fig_missingness_cards.png")
plt.close(fig)

# -------- Theory figure: Ng's elongated contours (scaling) --------
def gd_path(cx, cy, lr, steps, start=(2.0, 2.4)):
    x, y = start
    xs, ys = [x], [y]
    for _ in range(steps):
        x, y = x - lr * 2 * cx * x, y - lr * 2 * cy * y
        xs.append(x)
        ys.append(y)
    return xs, ys


fig, axes = plt.subplots(1, 2, figsize=(7.6, 3.9))
grid = np.linspace(-3, 3, 320)
W1, W2 = np.meshgrid(grid, grid)
for ax, (cx, cy, lr, steps, title) in zip(axes, [
        (25, 1, 0.037, 34, "unscaled: a long skinny valley"),
        (2, 2, 0.2, 8, "scaled: a nearly round bowl")]):
    Z = cx * W1 ** 2 + cy * W2 ** 2
    ax.contour(W1, W2, Z, levels=np.geomspace(0.4, Z.max() * 0.8, 9),
               colors=pal["sky"], linewidths=1.1)
    xs, ys = gd_path(cx, cy, lr, steps)
    ax.plot(xs, ys, ":o", color=pal["navy"], lw=1.2, ms=2.8)
    ax.plot(xs[0], ys[0], "o", color=pal["navy"], ms=5.5, zorder=4)
    ax.annotate("start", (xs[0], ys[0]), textcoords="offset points",
                xytext=(5, 4), fontsize=8.5, fontweight="bold",
                color=pal["navy"])
    ax.plot(0, 0, "*", color=pal["blue"], ms=13)
    ax.annotate("best", (0, 0), textcoords="offset points",
                xytext=(7, -3), fontsize=8.5, fontweight="bold",
                color=pal["blue"])
    ax.set_title(title, fontsize=11)
    ax.set_xlabel("the model's dial for Fare (its weight)", fontsize=8.5)
    ax.set_ylabel("the model's dial for Age (its weight)", fontsize=8.5)
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_aspect("equal")
axes[0].text(-2.8, -2.8, "the walk zigzags\nand crawls", fontsize=8.5,
             color=pal["ink"], ha="left", va="bottom",
             bbox=dict(facecolor="white", edgecolor="none", pad=1.5))
axes[1].text(-2.8, -2.8, "straight to the\nbest (star)", fontsize=8.5,
             color=pal["ink"], ha="left", va="bottom",
             bbox=dict(facecolor="white", edgecolor="none", pad=1.5))
fig.tight_layout()
fig.savefig(f"{FIGS_T2}/fig_contours_scaling.png")
plt.close(fig)

# -------- Theory figure: MIT's XOR - features bend the space --------
fig, axes = plt.subplots(1, 2, figsize=(7.6, 3.8),
                         gridspec_kw={"width_ratios": [1.1, 1]})
ax = axes[0]
for m in ((-2.2, 0.4), (0.6, -1.4), (-0.8, 2.6)):  # failed candidate lines
    xs = np.array([-1.9, 1.9])
    ax.plot(xs, m[1] * xs + m[0] * 0.3, "--", color=pal["gray"],
            lw=1.0, alpha=0.45)
ax.scatter([1, -1], [1, -1], s=110, color=pal["blue"], marker="o",
           zorder=3, label="one class")
ax.scatter([1, -1], [-1, 1], s=110, color=pal["navy"], marker="s",
           zorder=3, label="the other")
ax.legend(loc="upper left", fontsize=8, frameon=False)
ax.set_title("raw axes x1, x2: no line works", fontsize=11)
ax.set_xlabel("x1", fontsize=9.5)
ax.set_ylabel("x2", fontsize=9.5)
ax.set_xlim(-1.9, 1.9)
ax.set_ylim(-1.9, 1.9)
ax.set_xticks([-1, 0, 1])
ax.set_yticks([-1, 0, 1])
ax = axes[1]
ax.axhline(0, color=pal["gray"], lw=1.2)
ax.scatter([1, 1], [0.18, -0.18], s=110, color=pal["blue"], marker="o", zorder=3)
ax.scatter([-1, -1], [0.18, -0.18], s=110, color=pal["navy"], marker="s", zorder=3)
ax.axvline(0, color=pal["blue"], lw=1.6, ls="--")
ax.text(0.06, 0.85, "one threshold\nseparates them", fontsize=8.5,
        color=pal["blue"], ha="left")
ax.text(-1.72, 1.02, "height is only for\nvisibility - all that counts\nis left vs right",
        fontsize=7, color=pal["gray"], ha="left", va="top")
ax.text(-1, -0.55, "(-1,+1) and (+1,-1)\nland at -1", fontsize=8,
        color=pal["gray"], ha="center")
ax.text(1, -0.55, "(+1,+1) and (-1,-1)\nland at +1", fontsize=8,
        color=pal["gray"], ha="center")
ax.set_title("engineered axis x1 * x2: separable", fontsize=11)
ax.set_xlabel("x1 * x2", fontsize=9.5)
ax.set_xlim(-1.8, 1.8)
ax.set_ylim(-1.0, 1.2)
ax.set_xticks([-1, 0, 1])
ax.set_yticks([])
ax.spines["left"].set_visible(False)
fig.tight_layout()
fig.savefig(f"{FIGS_T2}/fig_xor_features.png")
plt.close(fig)

# -------- Theory figure: representation - same points, two coordinate systems --------
rng = np.random.default_rng(3)
x_rep = np.linspace(-2.0, 2.0, 26)
y_rep = 1.0 + 0.8 * x_rep ** 2 + rng.normal(0, 0.35, x_rep.size)
fig, axes = plt.subplots(1, 2, figsize=(7.6, 3.8))
ax = axes[0]
ax.scatter(x_rep, y_rep, s=42, color=pal["blue"], zorder=3)
b1, b0 = np.polyfit(x_rep, y_rep, 1)
xs = np.linspace(-2.15, 2.15, 100)
ax.plot(xs, b0 + b1 * xs, "--", color=pal["navy"], lw=1.8)
ax.set_title("raw axes (x, y): no line fits", fontsize=11)
ax.set_xlabel("x", fontsize=9.5)
ax.set_ylabel("y", fontsize=9.5)
ax = axes[1]
z_rep = x_rep ** 2
ax.scatter(z_rep, y_rep, s=42, color=pal["blue"], zorder=3)
c1, c0 = np.polyfit(z_rep, y_rep, 1)
zs = np.linspace(-0.1, 4.5, 100)
ax.plot(zs, c0 + c1 * zs, "-", color=pal["navy"], lw=1.8)
ax.set_title("new axis $x^2$: the same points line up", fontsize=11)
ax.set_xlabel("$x^2$", fontsize=9.5)
ax.set_ylabel("y", fontsize=9.5)
ax.text(0.97, 0.06, "same 26 points -\nonly the axis changed",
        transform=ax.transAxes, ha="right", va="bottom", fontsize=8.5,
        color=pal["ink"],
        bbox=dict(facecolor="white", edgecolor="none", pad=1.5))
fig.tight_layout()
fig.savefig(f"{FIGS_T2}/fig_representation.png")
plt.close(fig)

# -------- Theory figure: basis expansion - a line learns curves --------
rng = np.random.default_rng(11)
x_be = np.sort(rng.uniform(0.2, 3.8, 34))
y_be = 0.6 + 2.2 * x_be - 0.52 * x_be ** 2 + rng.normal(0, 0.28, x_be.size)
fig, axes = plt.subplots(1, 2, figsize=(7.6, 3.8), sharey=True)
xs = np.linspace(0.05, 3.95, 120)
ax = axes[0]
ax.scatter(x_be, y_be, s=40, color=pal["blue"], zorder=3)
d1, d0 = np.polyfit(x_be, y_be, 1)
ax.plot(xs, d0 + d1 * xs, "--", color=pal["navy"], lw=1.8)
ax.set_title("model sees [x]: a line underfits", fontsize=11)
ax.set_xlabel("x", fontsize=9.5)
ax.set_ylabel("y", fontsize=9.5)
ax = axes[1]
ax.scatter(x_be, y_be, s=40, color=pal["blue"], zorder=3)
e2, e1, e0 = np.polyfit(x_be, y_be, 2)
ax.plot(xs, e0 + e1 * xs + e2 * xs ** 2, "-", color=pal["navy"], lw=1.8)
ax.set_title('model sees [x, $x^2$]: the "line" bends', fontsize=11)
ax.set_xlabel("x", fontsize=9.5)
ax.text(0.03, 0.05, "same fitter -\none extra column",
        transform=ax.transAxes, ha="left", va="bottom", fontsize=8.5,
        color=pal["ink"],
        bbox=dict(facecolor="white", edgecolor="none", pad=1.5))
fig.tight_layout()
fig.savefig(f"{FIGS_T2}/fig_basis_expansion.png")
plt.close(fig)

# -------- Theory figure: the encodings menu (three honest codes) --------
MENU = [
    ("ONE-HOT", "no real order", pal["blue"], "white",
     "Embarked: C / Q / S",
     ["C  ->  1 0 0", "Q  ->  0 1 0", "S  ->  0 0 1"],
     "every pair of categories\nsits at the same distance"),
    ("THERMOMETER", "real order, no spacing", pal["navy"], "white",
     "T-shirt: S < M < L",
     ["S  ->  1 0 0", "M  ->  1 1 0", "L  ->  1 1 1"],
     "keeps the ranking,\ninvents no distances"),
    ("KEEP NUMERIC", "a true quantity", pal["sky"], pal["navy"],
     "Age, Fare",
     ["standardize:", "(value - mean) / std", "std = typical spread"],
     "real distances kept,\nunits made comparable"),
]
fig, ax = plt.subplots(figsize=(7.3, 4.3))
for i, (name, sub, head_fc, head_tc, example, code, honest) in enumerate(MENU):
    x = 0.15 + i * 4.0
    ax.add_patch(Rectangle((x, 1.55), 3.7, 6.75, facecolor=pal["panel"],
                           edgecolor=pal["sky"], lw=1.2))
    ax.add_patch(Rectangle((x, 8.3), 3.7, 1.5, facecolor=head_fc))
    ax.text(x + 1.85, 9.25, name, ha="center", va="center", fontsize=12.5,
            fontweight="bold", color=head_tc)
    ax.text(x + 1.85, 8.62, sub, ha="center", va="center", fontsize=7.4,
            color=head_tc)
    ax.text(x + 0.2, 7.75, "EXAMPLE", fontsize=6.4, fontweight="bold",
            color=pal["gray"], va="top")
    ax.text(x + 0.2, 7.2, example, fontsize=8.6, color=pal["ink"], va="top")
    y = 6.25
    for line in code:
        ax.text(x + 0.45, y, line, fontsize=8.6, color=pal["navy"],
                va="top", family="monospace", fontweight="bold")
        y -= 0.75
    ax.text(x + 0.2, 3.55, "WHY IT IS HONEST", fontsize=6.4,
            fontweight="bold", color=pal["gray"], va="top")
    ax.text(x + 0.2, 3.0, honest, fontsize=8.4, color=pal["ink"], va="top")
ax.add_patch(Rectangle((0.15, 0.15), 11.7, 1.15, facecolor=pal["navy"]))
ax.text(6.0, 0.725, "never: arbitrary integer codes (C=1, Q=2, S=3)\n"
        "they invent an order and a distance the data never had",
        ha="center", va="center", fontsize=8.5, color="white",
        fontweight="bold", linespacing=1.5)
ax.set_xlim(0, 12.1)
ax.set_ylim(0, 9.9)
ax.axis("off")
fig.savefig(f"{FIGS_T2}/fig_encodings_menu.png")
plt.close(fig)

# --- Figure 5 (story): Frankie's row (top lane) + the Pipeline (bottom lane) ---
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

fig, ax = plt.subplots(figsize=(7.6, 4.8))
ax.text(0.15, 7.65, "TOP LANE - READ THIS FIRST: ONE PASSENGER'S ROW, "
        "LEFT TO RIGHT", fontsize=7.5,
        fontweight="bold", color=pal["gray"])
top = [("Raw row (PassengerId 166)",
        'Goldsmith, Master. Frank\nJohn William "Frankie"\n'
        "Pclass 3, male, Age 9, SibSp 0,\nParch 2, Fare 20.525, Embarked S",
        0.15, 3.05, "panel"),
       ("Engineered",
        'Title = "Master"\nFamilySize = 0 + 2 + 1 = 3\nIsAlone = 0',
        3.85, 2.65, "blue"),
       ("18 numbers the model sees",
        "Title_Master, Pclass_3,\nSex_male, Embarked_S = 1\n"
        "(9 other one-hot columns 0);\n5 numeric columns scaled",
        7.15, 2.85, "navy")]
for head, body, x, w, kind in top:
    fc = {"panel": pal["panel"], "blue": pal["blue"], "navy": pal["navy"]}[kind]
    txt = pal["navy"] if kind == "panel" else "white"
    body_c = pal["ink"] if kind == "panel" else "white"
    ax.add_patch(FancyBboxPatch((x, 5.0), w, 2.45, boxstyle="round,pad=0.05",
                                facecolor=fc, edgecolor=pal["blue"], lw=1.2))
    ax.text(x + w / 2, 7.12, head, ha="center", va="center",
            fontweight="bold", color=txt, fontsize=8.6)
    ax.text(x + w / 2, 6.02, body, ha="center", va="center",
            color=body_c, fontsize=7.4)
for x0, lab in ((3.28, "engineer\nnew columns"), (6.58, "impute, one-hot,\nscale")):
    ax.add_patch(FancyArrowPatch((x0, 6.2), (x0 + 0.5, 6.2),
                                 arrowstyle="-|>", mutation_scale=13,
                                 color=pal["blue"], lw=1.5))
    ax.text(x0 + 0.25, 4.78, lab, ha="center", va="top",
            color=pal["gray"], fontsize=6.8)
ax.text(0.15, 3.85, "BOTTOM LANE - THEN THIS: THE PIPELINE THAT DOES IT",
        fontsize=7.5, fontweight="bold", color=pal["gray"])
bot = [("Impute", 'median Age 28.5,\nEmbarked "S" (from train)', 0.55, 2.75),
       ("Encode", "one-hot:\n13 x 0/1 columns", 3.95, 2.55),
       ("Scale", "StandardScaler,\n5 numeric columns", 7.15, 2.55)]
for head, body, x, w in bot:
    ax.add_patch(FancyBboxPatch((x, 1.55), w, 1.85, boxstyle="round,pad=0.05",
                                facecolor=pal["panel"], edgecolor=pal["blue"],
                                lw=1.2))
    ax.text(x + w / 2, 3.05, head, ha="center", va="center",
            fontweight="bold", color=pal["navy"], fontsize=8.6)
    ax.text(x + w / 2, 2.25, body, ha="center", va="center",
            color=pal["ink"], fontsize=7.4)
for x0 in (3.42, 6.62):
    ax.add_patch(FancyArrowPatch((x0, 2.45), (x0 + 0.4, 2.45),
                                 arrowstyle="-|>", mutation_scale=13,
                                 color=pal["blue"], lw=1.5))
# aligned link: the top lane's second arrow is implemented by the bottom lane
ax.plot([6.83, 6.83], [4.35, 5.9], ls=":", color=pal["navy"], lw=1.3)
ax.plot([6.83, 5.0], [4.35, 3.62], ls=":", color=pal["navy"], lw=1.3)
ax.text(7.05, 4.18, "dotted link: this arrow\nIS the bottom lane",
        fontsize=6.5, color=pal["navy"], ha="left", va="top")
ax.plot([0.55, 0.55, 9.70, 9.70], [1.15, 0.92, 0.92, 1.15],
        color=pal["navy"], lw=1.2)
ax.text(5.125, 0.62, "one Pipeline + ColumnTransformer: fit on train only, "
        "transform train and test", ha="center", va="top", color=pal["navy"],
        fontsize=8.4, fontweight="bold")
ax.text(5.125, 0.02, "Name, Ticket, Cabin, PassengerId, IsAlone: not on the "
        "feature lists - dropped automatically", ha="center", va="top",
        color=pal["gray"], fontsize=7.4)
ax.set_xlim(0, 10.2)
ax.set_ylim(-0.5, 7.9)
ax.axis("off")
fig.savefig(f"{FIGS_STORY}/s2_pipeline_frankie.png")
plt.close(fig)

# -------- Theory figure (Ackoff 1989): data -> information -> knowledge -> wisdom --------
# The DIKW ladder filled with one real passenger (PassengerId 166). Illustration.
STEPS = [
    ("DATA", "symbols - no meaning yet", "the raw file",
     'Goldsmith, Master. Frank\nJohn William "Frankie"\n3, male, 9, 0, 2,\n20.525, S'),
    ("INFORMATION", "organized for a purpose", "TODAY - Session 2",
     "a 9-year-old boy, 3rd class,\ntravelling with 2 parents:\nTitle = Master\nFamilySize = 3"),
    ("KNOWLEDGE", "a rule for using information", "Session 3",
     "boys in small families\nsurvived far more often\nthan grown men"),
    ("WISDOM", "acting well on knowledge", "Session 4",
     "predict for a new\npassenger, then decide\nwhat to do"),
]
fig, ax = plt.subplots(figsize=(7.6, 4.8))
bw, bh = 1.72, 1.95
for i, (name, plain, when, filling) in enumerate(STEPS):
    x = 0.12 + i * 1.92
    y0 = 0.25 + i * 0.62
    today = i == 1
    fc = pal["blue"] if today else pal["panel"]
    ec = pal["blue"] if today else pal["sky"]
    hc = "white" if today else pal["navy"]
    pc = pal["sky"] if today else pal["gray"]
    bc = "white" if today else pal["ink"]
    ax.add_patch(FancyBboxPatch((x, y0), bw, bh, boxstyle="round,pad=0.04",
                                facecolor=fc, edgecolor=ec, lw=1.5))
    ax.text(x + bw / 2, y0 + bh - 0.27, name, ha="center", va="center",
            fontsize=10.5, fontweight="bold", color=hc)
    ax.text(x + bw / 2, y0 + bh - 0.55, plain, ha="center", va="center",
            fontsize=7.4, color=pc, style="italic")
    ax.text(x + bw / 2, y0 + 0.62, filling, ha="center", va="center",
            fontsize=7.8, color=bc, linespacing=1.3)
    ax.text(x + bw / 2, y0 - 0.14, when, ha="center", va="top", fontsize=8,
            fontweight="bold", color=pal["blue"] if today else pal["gray"])
    if i:
        ax.add_patch(FancyArrowPatch((x - 0.19, y0 - 0.62 + bh / 2),
                                     (x + 0.01, y0 + bh / 2),
                                     arrowstyle="-|>", mutation_scale=14,
                                     color=pal["blue"], lw=1.6))
ax.set_xlim(0, 7.8)
ax.set_ylim(0, 4.3)
ax.axis("off")
fig.savefig(f"{FIGS_T2}/fig_dikw_titanic.png")
plt.close(fig)

# -------- Theory figure (Sweeney 2000): quasi-identifiers on the manifest --------
# Counts computed on the course dataset (891 rows; 714 with a known Age):
# share of passengers pinned down to exactly one person by the listed columns.
REID_LABELS = ["class + sex", "+ age", "+ port of boarding", "+ fare paid"]
REID_SHARES = [0.0, 16.9, 31.1, 86.3]
fig, ax = plt.subplots(figsize=(7.6, 4.8))
reid_bars = ax.barh(REID_LABELS, REID_SHARES,
                    color=[pal["sky"], pal["sky"], pal["sky"], pal["blue"]])
for bar, v in zip(reid_bars, REID_SHARES):
    ax.text(v + 1.5, bar.get_y() + bar.get_height() / 2, f"{v:.1f}%",
            va="center", fontweight="bold", color=pal["ink"])
ax.invert_yaxis()
ax.set_xlim(0, 100)
ax.set_xlabel("share of passengers these columns pin down to exactly ONE "
              "person\n(the 714 passengers with a known age)", fontsize=10)
ax.set_title("No name, no ID - yet four ordinary columns single out most "
             "passengers", fontsize=12.5)
ax.text(36, 1.5, "Florence: 1st class, female, 38, boarded\nCherbourg - 2 "
        "passengers match.\nAdd her fare: she is the only one.",
        va="center", fontsize=9.2, color=pal["navy"], linespacing=1.3,
        bbox=dict(boxstyle="round,pad=0.4", facecolor=pal["panel"],
                  edgecolor=pal["sky"]))
fig.tight_layout()
fig.savefig(f"{FIGS_T2}/fig_reidentify.png")
plt.close(fig)

# ================================ DECK ================================
prs = ds.new_deck()

# Slide 1 - title
ds.title_slide(
    prs, "Session 2 of 4",
    "Data Preparation and Feature Engineering",
    "The manifest - the ship's passenger list - is messy: 177 missing "
    "ages, titles buried inside names, scales that disagree. Today we "
    "turn passengers into signals - without cheating.")

# Slide 2 - REPLACED: Frame C introduction (FIT | CHOOSE | REPORT bar)
s = ds.image_slide(
    prs,
    "One bar rules the module: FIT on train, CHOOSE by validation, "
    "REPORT on test - once",
    f"{FIGS_T2}/frame_c_bar_s2.png",
    kicker="Recap and roadmap",
    bullets=[
        "Last time: 891 passenger rows, split 80/20 - 712 train, 179 "
        "test (stratified, seed 42)",
        "How to read the bar: pale band = the 712 training rows; dark "
        "block = the 179 test rows, sealed",
        "The hatched stripe sits INSIDE training: validation - practice "
        "tests, built in Session 3",
        "Exam terms: study on train, pick your strategy with practice "
        "tests, sit the final once",
        "Anything learned FROM data - a fill-in value, a scaler - is "
        "part of the model",
        "So it must never see the test rows. Today lives entirely in "
        "the FIT band",
    ],
    caption="The bar every session returns to. Harvard CS109A / ISLR ch. 5 "
            "(model selection vs model assessment).")
notes(s, "Teach the bar itself before the rules: the pale band is the 712 "
         "training rows, the dark block is the 179 test rows kept sealed, "
         "and the hatched stripe inside training is validation - the "
         "practice tests we build in Session 3. Then the exam metaphor: "
         "study on train, pick your strategy with practice tests, sit the "
         "final once. Today never leaves the study material - every value "
         "we compute comes from the 712 training rows. Ask the class: why "
         "does even an average count as 'part of the model'?")

# Slide 3 - Part 1 divider: the theory chapter
ds.section_slide(
    prs, "01", "Part 1 - The theory: turning passengers into rows of numbers",
    "A row of numbers is called a vector. Before any code: the model only "
    "sees the numbers you construct - gaps, encodings, scales, and new "
    "columns are design decisions, not chores")

# NEW - Ackoff's ladder: preparation is the step from data to information
s = ds.image_slide(
    prs,
    "Preparation is the step from data to information - the model cannot "
    "take it for you",
    f"{FIGS_T2}/fig_dikw_titanic.png",
    kicker="Data, information, knowledge",
    bullets=[
        "How to read it: four rising boxes, one real passenger climbing "
        "them; the blue box is today",
        "Data, in plain words: symbols with no meaning yet - '3, male, 9, "
        "0, 2, 20.525, S'",
        "Information: the same symbols organized for a purpose - a "
        "9-year-old boy travelling with two parents",
        "Knowledge: a rule for using information - Session 3 learns it; "
        "wisdom: acting well on it - Session 4 ships it",
        "Kitchen terms: raw ingredients, prepared ingredients, the recipe, "
        "the meal served",
        "Today's whole job is step one to step two - and no algorithm "
        "does it for you",
    ],
    caption="Illustration - the data-information-knowledge-wisdom ladder. "
            "Ackoff, 'From Data to Wisdom', Journal of Applied Systems "
            "Analysis, 1989. Passenger: PassengerId 166, fields verbatim.")
notes(s, "Start from the raw row and read it aloud as symbols: 3, male, 9, "
         "0, 2, 20.525, S - it means nothing until you know what each "
         "position is. Organize it for a purpose and it becomes "
         "information: a nine-year-old boy in third class travelling with "
         "two parents - Title Master, FamilySize 3. A rule for using that "
         "information is knowledge - Session 3's model learns it; acting "
         "well on the rule is wisdom - Session 4's app. Kitchen line: raw "
         "ingredients, prepared ingredients, recipe, meal. Ask the class: "
         "which step did M1's dashboards stop at?")

# NEW - Why preparation deserves the time: data cascades
s = ds.big_number_slide(
    prs,
    "Skimping on this step is the most common way real ML projects fail",
    "92%",
    "of AI practitioners in a Google study had lived through at least one "
    "'data cascade' - a data flaw that surfaced late, as a model failure",
    foot="In plain words: a cascade is a small data problem that grows as it "
         "flows downstream - a typo in a recipe's first line. 53 "
         "practitioners interviewed; 45% reported two or more per project. "
         "Sambasivan et al., 'Everyone wants to do the model work, not the "
         "data work', CHI 2021. Time budget: loading and cleaning data take "
         "about 45% of a data scientist's day (Anaconda, State of Data "
         "Science 2020).",
    kicker="Why this session exists")
notes(s, "The number first: 92 out of 100 practitioners in a Google Research "
         "study had lived through a data cascade - a small data flaw that "
         "surfaced late and expensively as a model failure. The paper's "
         "title says the rest: everyone wants to do the model work, not the "
         "data work. Then the time budget: about 45% of a data scientist's "
         "day goes to loading and cleaning. Today is that 45%. Ask the "
         "class: name one gap in the manifest that could become a cascade "
         "if we ignored it.")

# NEW - Data quality has dimensions (Wang & Strong)
s = ds.table_slide(
    prs,
    "Data quality is more than accuracy - each manifest column fails a "
    "different dimension",
    ["Dimension (in plain words)", "On the manifest", "What we do today"],
    [
        ["Completeness - are the values there?",
         "Age blank in 177 rows; Cabin blank in 687 of 891",
         "Impute Age from train; drop Cabin"],
        ["Accuracy - are they right?",
         "18 ages end in .5 - the file's mark for an estimated age; 15 "
         "fares read 0.00",
         "Keep, but remember Session 1's noise floor"],
        ["Consistent representation - same thing, same form?",
         "Embarked is letters C/Q/S, Sex is words, Fare is numbers",
         "One-hot the letters, scale the numbers"],
        ["Relevance - does it bear on the question?",
         "PassengerId and Ticket carry no signal about survival",
         "Drop them"],
        ["Appropriate amount - the right level of detail?",
         "Name is raw text; 14 different titles hide inside it",
         "Mine Name for Title, group 14 into 5 buckets"],
        ["Believability - would you trust it?",
         "A fare of 0.00 - free ticket, staff perk, or missing?",
         "Flag it and ask someone who knows the domain"],
    ],
    kicker="Data quality: the theory",
    note="Dimensions from Wang & Strong, 'Beyond Accuracy: What Data Quality "
         "Means to Data Consumers', J. of Management Information Systems, "
         "1996. Counts computed on the full 891-row file; the .5 rule is the "
         "dataset's own convention for estimated ages.",
    col_widths=[0.30, 0.40, 0.30])
notes(s, "Kitchen line: inspect the ingredients before you cook. Read one "
         "row at a time and say the dimension in plain words first - is the "
         "value there, is it right, is it in the same form, does it matter, "
         "is it the right level of detail, would you trust it. Every "
         "Titanic column fails a different one, and every failure maps to "
         "one of today's moves. Ask the class: 15 fares read 0.00 - which "
         "dimension is that, and what would you do before imputing?")

# NEW - Anonymous is not unidentifiable (Sweeney)
s = ds.image_slide(
    prs,
    "Dropping the Name column is not anonymisation - four ordinary "
    "columns single out 86% of passengers",
    f"{FIGS_T2}/fig_reidentify.png",
    kicker="Data you must handle with care",
    bullets=[
        "We drop Name and PassengerId today - the model does not need "
        "them",
        "How to read it: each bar adds one column; bar length = share of "
        "passengers left standing alone",
        "Class + sex + age + port already pins down 31%; add the fare "
        "paid, 86%",
        "In plain words: quasi-identifiers are columns that are not names "
        "but together point at one person",
        "Sweeney (2000): ZIP code + birth date + sex uniquely identify 87% "
        "of Americans",
        "The Titanic is public history; your customers' tables are not - "
        "keep only the columns the question needs",
    ],
    caption="Counts computed on the course dataset (714 passengers with a "
            "known age). Sweeney, 'Simple Demographics Often Identify People "
            "Uniquely', Carnegie Mellon, 2000.")
notes(s, "Ground it in a name they know: Florence Cumings. Drop her name and "
         "ID, keep class, sex, age and port - only two passengers on the "
         "whole ship match her; add her fare and she is alone. Read the "
         "bars: each adds one ordinary column, and the share of passengers "
         "standing alone jumps from 0 to 86%. Sweeney showed the same for "
         "Americans with ZIP code, birth date and sex. The Titanic is public "
         "history, so this is a lesson, not a leak; their future customer "
         "tables are different. Ask the class: which of today's engineered "
         "columns makes a passenger easier to single out?")

# Slide 4 - NEW: feature representation is a first-class decision (MIT)
s = ds.image_slide(
    prs,
    "The model only sees the numbers you construct - representation "
    "comes first",
    f"{FIGS_T2}/fig_representation.png",
    kicker="Feature representation",
    bullets=[
        "The model never meets a passenger - it meets a row of numbers "
        "built to stand for one",
        "That row is the representation: the numbers you choose to "
        "describe each passenger",
        "How to read the charts: each dot is one example; the line is "
        "the best straight line",
        "Left: plotted against x, no straight line fits the dots",
        "Right: the SAME dots against x squared - now they line up",
        "The dots never move; only the axis changes - good columns beat "
        "fancy algorithms",
    ],
    caption="Same 26 synthetic dots in both panels; only the horizontal "
            'axis changes. MIT 6.390, ch. 5 "Feature representation".')
notes(s, "The model never meets Frankie - it meets the row of numbers you "
         "built to stand for him, and that row is its entire world. Point "
         "at the two panels: the dots never move, only the axis changes, "
         "and suddenly a straight line fits. That is the whole chapter in "
         "one picture. Ask the class: if the model only sees the numbers, "
         "who is responsible when it sees the world wrongly?")

# Slide 5 - the three requirements (promoted into the theory chapter)
s = ds.table_slide(
    prs,
    "Models are picky eaters: they want numbers only, no gaps, sensible scales",
    ["The model demands", "Raw Titanic reality", "The fix"],
    [
        ["Numbers only",
         'Sex is "male"/"female"; Embarked is C/Q/S; Name is free text',
         "Encoding: turn categories into 0/1 columns"],
        ["No gaps",
         "Age missing in 177 of 891 rows (19.9%); Embarked missing in 2",
         "Imputation: fill gaps with a value learned from train"],
        ["Sensible scales",
         "Age runs 0 to 80 while Fare runs 0 to 512 - Fare shouts, "
         "Age whispers",
         "Scaling: put features on a comparable ruler"],
    ],
    kicker="Why preparation exists",
    note="Why so picky? A model sees each passenger as a point in space, one "
         "axis per column - words and gaps have no place on an axis. Impute = "
         "fill a gap with a learned substitute. Counts: Module 2, notebook 1.",
    col_widths=[0.22, 0.44, 0.34])
notes(s, "The kitchen line: models are picky eaters - they only eat "
         "numbers, refuse plates with gaps, and complain when portions are "
         "on wildly different scales. Each row of the table is one "
         "complaint and its fix: encoding, imputation, scaling. Ask the "
         "class: what would the model make of the word 'Southampton' if we "
         "fed it raw?")

# Slide 6 - MCAR / MAR / MNAR three-card theory slide
s = ds.image_slide(
    prs,
    "Missing values come in three species - and the data cannot tell you which",
    f"{FIGS_T2}/fig_missingness_cards.png",
    kicker="Missing values: the theory",
    bullets=[
        "Our table is missing 177 of 891 ages. WHY a hole exists matters",
        "How to read: one card per species - meaning, a Titanic story, "
        "the risk, the fix",
        "MCAR: holes poked blindfolded - nothing decides where they land",
        "MAR: holes cluster by ANOTHER column you can see (class, say)",
        "MNAR: the hidden value itself caused the hole - the nastiest kind",
        "The data alone cannot tell you which. Discuss: which are our 177?",
    ],
    caption="Each Titanic story is plausible, not asserted. Harvard CS109A, "
            "Missing Data lecture (2020, L19).")
notes(s, "Ground it first: 177 of our 891 ages are simply blank - the "
         "question is WHY. Walk the cards left to right; each card gives "
         "the meaning, a plausible Titanic story, the risk, the fix. "
         "Harvard's device: MCAR is poking holes in the table blindfolded - "
         "annoying but fair. MAR holes cluster where another visible column "
         "points, so ignoring them biases you. MNAR holes are caused by the "
         "very value that is missing - no fill can see that. The deciding "
         "information is missing by definition, so the table cannot settle "
         "it. Run the discussion: which species are our 177 missing ages? "
         "Any argued answer beats a certain one.")

# Slide 7 - NEW: the encodings menu (MIT)
s = ds.image_slide(
    prs,
    "Three honest encodings: one-hot labels, thermometer ranks, "
    "standardized numbers",
    f"{FIGS_T2}/fig_encodings_menu.png",
    kicker="Encoding: the theory",
    bullets=[
        "The manifest says Embarked = S. The model needs numbers, not "
        "letters",
        "Encoding = translating categories into numbers, without "
        "inventing facts",
        "How to read: one card per code - when it applies, the "
        "translation, why it is honest",
        "One-hot (no real order): each category gets its own 0/1 column",
        "Thermometer (order, no spacing): keeps the ranking S < M < L, "
        "nothing more",
        "Keep-numeric (true quantities): Age and Fare stay numbers",
        "The navy strip is the trap: C=1, Q=2, S=3 claims S is three "
        "times C",
    ],
    caption="Illustration - the encoding menu. MIT 6.390, ch. 5, sec. 5.3 "
            '"Hand-constructing features".')
notes(s, "Start concrete: the manifest literally says 'S' in the Embarked "
         "column, and the model cannot eat letters. Walk the three cards "
         "left to right, then the navy strip at the bottom - the forbidden "
         "move. Frame encoding as translation into the only language the "
         "model speaks - numbers - and honesty as translating without "
         "adding claims. Coding red=1, blue=2 quietly tells the model blue is "
         "twice red; it will believe you. One-hot claims nothing, the "
         "thermometer claims only order, keeping a real number claims real "
         "distances. Ask the class: is T-shirt size S/M/L more like "
         "Embarked or more like Age?")

# Slide 8 - Ng's elongated contours - why scaling helps training
s = ds.image_slide(
    prs,
    "Scaling does not change the destination - it straightens the road to it",
    f"{FIGS_T2}/fig_contours_scaling.png",
    kicker="Feature scaling: the why",
    bullets=[
        "Fare runs 0 to 512, Age 0 to 80 - to the model, Fare shouts "
        "and Age whispers",
        "Each picture is a map of the search for the best settings: "
        "every ring = settings scoring the same; the star = the best; "
        "the dots = training's steps",
        "Unscaled: the map is a long skinny valley - the search zigzags "
        "and crawls",
        "Scaled: the map turns round - the search walks straight to "
        "the star",
        "Scaling never changes the answer - only how fast and surely "
        "you find it",
    ],
    caption="Illustration - a schematic map of the search, not course "
            "data. Ng, Machine Learning Specialization C1W2.")
notes(s, "Ground it first: Fare runs 0-512 and Age 0-80, so to the model "
         "Fare shouts and Age whispers. Then teach the picture before the "
         "moral: each panel is a map of the search for the best settings - "
         "one ring joins settings that score the same, the star is the "
         "best setting, the dotted trail is training's steps from the "
         "'start' marker. Unscaled, the map is a skinny valley: the walk "
         "zigzags and crawls. Scaled, it is a round bowl: straight to the "
         "star. The destination never moves - only the trip changes. Ask "
         "the class: which of our columns stretches the valley worst?")

# Slide 9 - MIT's XOR - new features bend the space
s = ds.image_slide(
    prs,
    "One engineered column, x1 times x2, splits four points no line "
    "ever could",
    f"{FIGS_T2}/fig_xor_features.png",
    kicker="Feature engineering: the geometry",
    bullets=[
        "Four points, two facts each (x1, x2) - every value +1 or -1",
        "How to read: blue circles = one class, dark squares = the other",
        "Left: the gray dashed lines are failed tries - no straight line "
        "can split circles from squares",
        "Right: re-plot each point by ONE new number, x1 times x2, on a "
        "number line",
        "Circles land at +1, squares at -1 - one cut separates them",
        "We never changed the model - we changed the space it looks at",
        "Part 2 repeats the move on the manifest: Title and FamilySize",
    ],
    caption="Illustration - the classic XOR construction. MIT 6.390, ch. 5.")
notes(s, "Walk the picture slowly: four corner points - blue circles are "
         "one class, dark squares the other - and diagonal corners share "
         "a class, so no straight line will ever split them; let "
         "students try (the gray dashed lines are failed tries). Then "
         "multiply the two coordinates into one new "
         "column and the four points sort themselves onto two spots a "
         "single threshold separates. We never touched the model - we "
         "changed what it looks at. Ask the class: did the model get "
         "smarter, or did the world get simpler?")

# Slide 10 - NEW: basis expansion (Harvard/ISLR)
s = ds.image_slide(
    prs,
    "Add an x-squared column and the straight line learns to curve",
    f"{FIGS_T2}/fig_basis_expansion.png",
    kicker="Basis expansion",
    bullets=[
        "These dots rise, then fall - one straight line cannot follow "
        "that",
        "How to read: the same dots in both panels; left, the model "
        "sees only the column x",
        "Right: append ONE new column, x squared - the same fitter "
        "now bends",
        "Kitchen terms: same cook, one extra ingredient - now it can "
        "cook curves",
        "This is basis expansion: add powers and products of columns "
        "you already have (XOR's x1 times x2 is the same family)",
        "The price: every extra column can chase noise - Session 3 "
        "counts that cost",
    ],
    caption="Illustration - one fitter, two feature tables, synthetic "
            'points. Harvard CS109A, Lecture 4 "Polynomial Regression" '
            "/ ISLR.")
notes(s, "Point at the dots first: they rise and then fall, and a ruler "
         "cannot follow that - the dashed line on the left proves it. "
         "Then the kitchen: the line is the cook, the columns are the "
         "ingredients. Hand the cook x squared as one more ingredient and "
         "the same recipe starts producing curves - no new cook, no new "
         "technique. But every extra ingredient is also freedom to cook "
         "nonsense that matches noise; Session 3 puts a price on that. Ask "
         "the class: what column would you hand the line so it could bend "
         "twice?")

# Slide 11 - Part 2 divider: the practice chapter
ds.section_slide(
    prs, "02", "Part 2 - Practice: preparing the manifest",
    "The theory applied to 891 passengers: fit on the 712 training rows "
    "only, engineer Title and FamilySize, and wire every step into one "
    "leak-proof pipeline")

# Slide 12 - imputation the ML way
s = ds.image_slide(
    prs,
    "Fill gaps with a value learned from train only - the test median "
    "stays untouched",
    f"{FIGS}/fig_impute.png",
    kicker="Missing values the ML way",
    bullets=[
        "Age is blank in 137 train rows and 40 test rows (left chart)",
        "Median = the middle value. The 712 train rows say: median "
        "Age 28.5",
        "That one number, 28.5, fills every gap - in train AND in test",
        'The test-only median, 27.0, is "exactly the information we '
        'refuse to use"',
        'In code: SimpleImputer(strategy="median") - fit on train, '
        "transform both",
        'Embarked, 2 gaps: strategy="most_frequent" learns "S" from train',
    ],
    caption="Left: rows with Age blank. Right: each half's median - only "
            "train's is ever used. Module 2, notebook 1.")
notes(s, "Ground it in the chart: 137 blank ages in train, 40 in test. The "
         "median is just the middle value, and only the 712 training rows "
         "may vote: they say 28.5, and that one number fills every gap in "
         "both halves. The test rows' own median, 27.0, is exactly the "
         "information we refuse to use - the test set is the future, and "
         "you cannot compute statistics on data you have not seen yet. "
         "Same rule for Embarked's 2 gaps: most_frequent learns 'S' from "
         "train. Ask the class: what would be wrong with filling test "
         "gaps with 27.0?")

# Slide 13 - MERGED: one-hot by hand + the get_dummies trap
s13 = ds.table_slide(
    prs,
    "One-hot gives each category its own 0/1 column - and the training set "
    "fixes the columns once",
    ["Passenger", "Embarked", "Embarked_C", "Embarked_Q", "Embarked_S"],
    [
        ["1", "S", "0", "0", "1"],
        ["2", "C", "1", "0", "0"],
        ["3", "Q", "0", "1", "0"],
        ["4", "S", "0", "0", "1"],
    ],
    kicker="Encoding categorical features",
    note="Part 1's menu chose for us: Embarked has no real order, so "
         "one-hot - never C=1, Q=2, S=3. Module 2, notebook 1.",
    col_widths=[0.18, 0.22, 0.20, 0.20, 0.20])
s13b = ds._box(s13, ds.MARGIN, Inches(4.68), ds.CONTENT_W, Inches(1.6))
for i, line in enumerate([
        "Read one row: Passenger 2 boarded at C, so Embarked_C is 1 and "
        "the other two are 0 - exactly one column is 'hot' per row",
        "fit(train) fixes the column layout once; transform(test) always "
        "reproduces those same columns",
        'handle_unknown="ignore": an unseen category becomes all zeros - '
        "no crash",
        "pd.get_dummies instead builds columns from whatever it happens to "
        "see - train and test drift apart (the notebook demo misaligns them)"]):
    ds._para(s13b.text_frame, line, 13.5, ds.INK, first=(i == 0), bullet=True,
             space_after=6)
notes(s13, "Read one row of the table aloud: Passenger 2 boarded at C, so "
           "the C column gets the 1 and the others get 0 - that lone 1 is "
           "why it is called one-hot. No order, no distance is invented; "
           "Part 1's menu chose this because Embarked has no real order. "
           "Then the trap: pd.get_dummies rebuilds columns from whatever "
           "data it sees, so train and test can drift apart; the sklearn "
           "encoder fixes the layout at fit(train) and reproduces it "
           "forever. Ask the class: what should happen to a port the "
           "training set never saw?")

# Slide 14 - scaling on the real manifest
s = ds.image_slide(
    prs,
    "Scaling re-labels the ruler: distance models care, trees never notice",
    f"{FIGS}/fig_scales.png",
    kicker="Feature scaling",
    bullets=[
        "Look right: Fare runs 0 to 512, Age 0 to 80 - Fare shouts, "
        "Age whispers",
        "Distance and equation models (KNN, Logistic Regression) hear "
        "only the shouting axis: raw Fare gaps drown out Age gaps",
        'Trees only ask yes/no questions - "is Fare > 30?" - so scale '
        "never matters to them",
        "StandardScaler re-expresses values as 'how far from average': "
        "Fare becomes mean -0.00, std 1.00",
        "MinMaxScaler instead squeezes Fare into 0.0 to 1.0",
        "The histogram's shape never changes - scaling only re-labels "
        "the ruler",
    ],
    caption="Each bar = one column's raw range in the training set. "
            "Module 2, notebook 1.")
notes(s, "Same grounding as Part 1: Fare shouts (0-512), Age whispers "
         "(0-80) - the chart makes the mismatch physical. Models that "
         "measure distances or weigh sums (KNN, Logistic Regression) hear "
         "only the loud axis; trees just ask yes/no questions like 'is "
         "Fare > 30?', so they never notice. StandardScaler re-expresses "
         "each value as 'how far from average' - the mean and standard "
         "deviation from M1 are exactly the two numbers it learns from "
         "train; MinMaxScaler squeezes "
         "into 0-1. Either way the histogram's shape is untouched - only "
         "the ruler's labels change. Ask the class: which of our six "
         "models from the next session will care?")

# Slide 15 - Title from Name
s = ds.image_slide(
    prs,
    "Title packs sex, age, and status into one column - and finds the boys",
    f"{FIGS}/fig_titles.png",
    kicker="Feature engineering: Title",
    bullets=[
        'Every Name hides a title: "Cumings, Mrs. John Bradley" '
        "contains Mrs",
        "Name is unstructured text inside a structured table - free "
        "text has no column until you build one",
        'A regex (a text-search pattern), " ([A-Za-z]+)\\.", grabs the '
        "word before the period",
        "14 raw titles group into 5 buckets: Mr, Miss, Mrs, Master, Rare "
        "- the chart counts them",
        'Frankie, age 9, is a "Master" - the old title for a young boy',
        "Sex alone calls Owen (22) and Frankie (9) both male - Title "
        "tells the man from the boy",
        "Signal check: Mrs and Miss survived far more often than Mr; "
        "Master sits well above adult men",
    ],
    caption="Grouped Title counts, training set (712 rows). Master "
            "highlighted: the boys the Sex column hides. Module 2, "
            "notebook 1.")
notes(s, "Start from a real name: 'Cumings, Mrs. John Bradley' - the title "
         "Mrs is sitting right there in the text. The regex just grabs "
         "the word before the period; 14 raw titles collapse into five "
         "buckets. The payoff is Master, the old-fashioned title for a "
         "young boy: to the Sex column, Owen (22) and Frankie (9) are "
         "both simply 'male', but Title tells the man from the boy - "
         "which mattered a lot on a sinking ship. Signal check before "
         "keeping any new column: Mrs and Miss survived far more often "
         "than Mr, and Master sits well above adult men.")

# Slide 16 - FamilySize and IsAlone
s = ds.image_slide(
    prs,
    "FamilySize = SibSp + Parch + 1 - small families of 2 to 4 fared best",
    f"{FIGS}/fig_family.png",
    kicker="Feature engineering: FamilySize",
    bullets=[
        "FamilySize = SibSp (siblings/spouses) + Parch (parents/children) "
        "+ 1: you count yourself",
        "Owen: 1 + 0 + himself = 2. Frankie: 0 + 2 + himself = 3",
        "IsAlone = 1 when FamilySize is 1 - true for 434 of 712 training "
        "passengers",
        "Chart: one bar per family size, counting passengers; the blue "
        "bars are sizes 2 to 4",
        "Those small families fared best; solo travelers and very large "
        "families fared badly",
        "More features is not automatically better - always check for "
        "signal",
    ],
    caption="FamilySize distribution, training set. Blue bars: the sizes "
            "that fared best. Module 2, notebook 1.")
notes(s, "One addition, one new column: siblings/spouses plus parents/"
         "children plus yourself. Work the two passengers: Owen has a "
         "sibling aboard, so 1+0+1 = 2; Frankie has two parents aboard, "
         "so 0+2+1 = 3. IsAlone flags the 434 solo travelers - more than "
         "half the training set. On the chart, the blue bars (sizes 2-4) "
         "are the families that fared best; solo travelers and the very "
         "large families fared badly. Close with the discipline: a new "
         "column earns its place only if it carries signal.")

# Slide 17 - Ng's area feature
s = ds.two_col_slide(
    prs,
    "Multiply two weak columns and you can get one strong feature - Ng's area example",
    ("Andrew Ng - Machine Learning Specialization",
     [
         "A house lot has a street-facing width (x1, the 'frontage') "
         "and a depth (x2)",
         "New feature: x3 = x1 * x2 = the lot's area",
         "Exactly Part 1's x1 * x2 move, in Ng's own example",
         "Area often predicts price better than width or depth alone",
         "One multiplication - no new data collected",
     ]),
    ("The same move on the Titanic",
     [
         "SibSp + Parch + 1 = FamilySize: one column beats two",
         "Name -> Title: sex, age, and status compressed into one",
         "Master = young boy: a signal Sex and Age only hint at",
         "Good features often improve a model more than a fancier algorithm",
     ]),
    kicker="Feature engineering, attributed",
    note="Ng's rule for inventing features, paraphrased: use knowledge or "
         "intuition about the problem to design new columns. Ng, Machine "
         "Learning Specialization.")
notes(s, "Ng's own example, in plain terms: a house lot has a width facing "
         "the street (the 'frontage') and a depth. Multiply them and you "
         "get the area - one new column, often more predictive than "
         "either raw one, and no new data collected. Then mirror it on "
         "the Titanic: SibSp + Parch + 1 became FamilySize, and Name "
         "became Title. Same rule both times: use knowledge about the "
         "problem to design new columns. Ask the class: what other "
         "Titanic column pair might multiply into something useful?")

# Slide 18 - MERGED: Frankie's row + the pipeline, one aligned diagram
s = ds.image_slide(
    prs,
    "Frankie walks the pipeline: raw row in, 18 leak-proof numbers out",
    f"{FIGS_STORY}/s2_pipeline_frankie.png",
    kicker="Putting it all together",
    bullets=[
        "Read the top lane first: Frankie's one raw row becomes 18 "
        "numbers, left to right",
        "Then the bottom lane: the machinery - impute gaps, encode "
        "categories, scale numbers",
        "The dotted link: the top lane's last arrow IS the bottom lane",
        "fit_transform(train) learns every step's values; transform(test) "
        "only applies them",
        "Engineers call this shape ETL - extract, transform, load; one "
        "Pipeline, no leakage, nothing forgotten",
        "IsAlone is computed but not on the feature lists - the "
        "ColumnTransformer drops it",
    ],
    caption="PassengerId 166, raw fields verbatim from the dataset. "
            "The canonical prep, reused verbatim by Sessions 3 and 4. "
            "Module 2, notebook 1.")
notes(s, "Give the reading order out loud: top lane first - one real boy, "
         "Frankie, PassengerId 166, raw row on the left, 18 numbers on "
         "the right. Then the bottom lane: the three machines that do it "
         "- impute, encode, scale - and the dotted link showing the top "
         "lane's last arrow IS that machinery. Hammer the verbs once "
         "more: fit_transform on train learns the values, transform on "
         "test only applies them. One Pipeline means no leakage, no "
         "forgotten steps, and the whole thing travels as a single "
         "object - Sessions 3 and 4 reuse it verbatim. Note IsAlone: "
         "computed, but dropped by the ColumnTransformer.")

# Slide 19 - close (two-part recap, absorbs the smoke-test number)
ds.close_slide(
    prs,
    "Two parts, three golden rules",
    [
        "Part 1: preparation turns data into information (Ackoff) - the "
        "model sees only the numbers you build",
        "Quality has dimensions; missing values have species (MCAR, MAR, "
        "MNAR) the data cannot tell apart",
        "Features change the geometry - Ng's area, MIT's XOR, Harvard's "
        "x-squared: good columns beat fancy models",
        "Anonymous is not unidentifiable: four ordinary columns single "
        "out 86% of passengers",
        "Part 2: split first (80/20, seed 42, stratified); fit on train "
        "only; Pipelines always",
        "Proof: Logistic Regression scores 0.838 on the 179 unseen "
        "passengers, 0.829 on train - no memorization gap",
        "Session 3 chooses the model. Practice now: notebook "
        "01-data-prep-and-feature-engineering in Colab",
    ])

out = "../m2-session-2-data-prep-and-feature-engineering.pptx"
ds.save_deck(prs, out, "M2 Session 2 - Data Prep and Feature Engineering")
print(f"Slides: {len(prs.slides)}")
print(f"Saved: {os.path.abspath(out)}")
